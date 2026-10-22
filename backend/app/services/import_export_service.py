import csv
import io
import uuid
import openpyxl
from datetime import datetime, timezone
from typing import Dict, Any, List, Optional
from sqlalchemy.orm import Session
from fastapi import HTTPException

from app.models.models import Customer, CustomerLoyaltyAccount, LoyaltyProgram, LoyaltyTransaction, Campaign, Tenant
from app.services.customer_service import CustomerService
from app.services.audit_service import AuditService

# In-memory token store for previewing uploaded spreadsheets
_UPLOAD_CACHE: Dict[str, List[Dict[str, Any]]] = {}

class ImportExportService:
    @staticmethod
    def parse_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        """Parses CSV or Excel file and returns column headers and first 5 sample rows"""
        rows: List[Dict[str, Any]] = []

        if filename.endswith(".xlsx") or filename.endswith(".xls"):
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet = wb.active
            iter_rows = list(sheet.iter_rows(values_only=True))
            if not iter_rows:
                raise HTTPException(status_code=400, detail="Uploaded Excel sheet is empty")
            headers = [str(h).strip() if h is not None else f"Column_{i}" for i, h in enumerate(iter_rows[0])]
            for r in iter_rows[1:]:
                if not any(r):
                    continue
                row_dict = {}
                for idx, col_name in enumerate(headers):
                    val = r[idx] if idx < len(r) else None
                    row_dict[col_name] = str(val).strip() if val is not None else ""
                rows.append(row_dict)
        else:
            # CSV file
            content_str = file_bytes.decode("utf-8-sig", errors="replace")
            reader = csv.DictReader(io.StringIO(content_str))
            headers = reader.fieldnames or []
            for r in reader:
                rows.append({k.strip(): (v.strip() if v else "") for k, v in r.items() if k})

        if not headers:
            raise HTTPException(status_code=400, detail="No columns detected in uploaded file")

        file_token = str(uuid.uuid4())
        _UPLOAD_CACHE[file_token] = rows

        return {
            "file_token": file_token,
            "filename": filename,
            "total_rows": len(rows),
            "headers": headers,
            "sample_rows": rows[:5]
        }

    @staticmethod
    def execute_customer_import(
        db: Session,
        tenant_id: str,
        file_token: str,
        mapping: Dict[str, str],
        user_email: Optional[str] = None
    ) -> Dict[str, Any]:
        rows = _UPLOAD_CACHE.get(file_token)
        if rows is None:
            raise HTTPException(status_code=400, detail="Import session expired or invalid. Please re-upload the file.")

        phone_key = mapping.get("phone_col")
        first_name_key = mapping.get("first_name_col")
        last_name_key = mapping.get("last_name_col")

        if not phone_key:
            raise HTTPException(status_code=400, detail="Phone number column mapping is required")

        imported = 0
        updated = 0
        skipped = 0
        invalid = 0
        errors = []

        for idx, row in enumerate(rows):
            raw_phone = row.get(phone_key, "")
            if not raw_phone:
                skipped += 1
                continue

            clean_phone = CustomerService.normalize_phone(raw_phone)
            # Basic validation
            digits_only = "".join(filter(str.isdigit, clean_phone))
            if len(digits_only) < 10:
                invalid += 1
                errors.append(f"Row {idx+1}: Invalid phone number '{raw_phone}'")
                continue

            existing = db.query(Customer).filter(
                Customer.tenant_id == tenant_id,
                Customer.phone_normalized == clean_phone
            ).first()

            first_name = row.get(first_name_key) if first_name_key else None
            last_name = row.get(last_name_key) if last_name_key else None

            if existing:
                if first_name and not existing.first_name:
                    existing.first_name = first_name
                if last_name and not existing.last_name:
                    existing.last_name = last_name
                updated += 1
            else:
                cust = Customer(
                    tenant_id=tenant_id,
                    phone=raw_phone,
                    phone_normalized=clean_phone,
                    first_name=first_name,
                    last_name=last_name,
                    source="EXCEL_IMPORT",
                    marketing_consent=True,
                    whatsapp_consent=True
                )
                db.add(cust)
                db.flush()

                # Create loyalty account with welcome bonus
                program = db.query(LoyaltyProgram).filter(LoyaltyProgram.tenant_id == tenant_id).first()
                init_points = program.welcome_bonus_points if program else 100
                acc = CustomerLoyaltyAccount(
                    tenant_id=tenant_id,
                    customer_id=cust.id,
                    points_balance=init_points,
                    lifetime_earned=init_points,
                    lifetime_redeemed=0,
                    current_tier="BRONZE"
                )
                db.add(acc)
                imported += 1

        db.commit()
        # Clean up cache
        _UPLOAD_CACHE.pop(file_token, None)

        AuditService.log(
            db=db,
            action="CUSTOMERS_IMPORTED",
            tenant_id=tenant_id,
            user_email=user_email,
            resource_type="CUSTOMERS",
            details={"imported": imported, "updated": updated, "skipped": skipped, "invalid": invalid}
        )

        return {
            "imported": imported,
            "updated": updated,
            "skipped": skipped,
            "invalid": invalid,
            "errors": errors[:10]  # sample of first 10 errors
        }

    @staticmethod
    def export_customers_csv(db: Session, tenant_id: str) -> str:
        customers = db.query(Customer).filter(Customer.tenant_id == tenant_id).all()
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "ID", "Phone", "First Name", "Last Name", "Email", "Gender",
            "Date of Birth", "Address", "Points Balance", "Tier", "Lifetime Spend",
            "Visits Count", "WhatsApp Consent", "Created At"
        ])
        for c in customers:
            acc = db.query(CustomerLoyaltyAccount).filter(CustomerLoyaltyAccount.customer_id == c.id).first()
            writer.writerow([
                c.id,
                c.phone,
                c.first_name or "",
                c.last_name or "",
                c.email or "",
                c.gender or "",
                c.date_of_birth or "",
                c.address or "",
                acc.points_balance if acc else 0,
                acc.current_tier if acc else "BRONZE",
                f"{c.lifetime_spend:.2f}",
                c.visits_count,
                "YES" if c.whatsapp_consent else "NO",
                c.created_at.strftime("%Y-%m-%d %H:%M:%S") if c.created_at else ""
            ])
        return output.getvalue()

    @staticmethod
    def export_loyalty_transactions_csv(db: Session, tenant_id: str) -> str:
        txns = db.query(LoyaltyTransaction).filter(
            LoyaltyTransaction.tenant_id == tenant_id
        ).order_by(LoyaltyTransaction.created_at.desc()).all()

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Transaction ID", "Customer Phone", "Type", "Points", "Balance Before", "Balance After", "Description", "Date"])
        for t in txns:
            cust = db.query(Customer).filter(Customer.id == t.customer_id).first()
            writer.writerow([
                t.id,
                cust.phone if cust else "",
                t.type,
                t.points,
                t.balance_before,
                t.balance_after,
                t.description or "",
                t.created_at.strftime("%Y-%m-%d %H:%M:%S") if t.created_at else ""
            ])
        return output.getvalue()

    @staticmethod
    def export_campaigns_csv(db: Session, tenant_id: str) -> str:
        campaigns = db.query(Campaign).filter(Campaign.tenant_id == tenant_id).order_by(Campaign.created_at.desc()).all()
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Campaign ID", "Name", "Status", "Recipients", "Delivered", "Read", "Failed", "Actual Cost (INR)", "Attributed Revenue (INR)", "Created At"])
        for c in campaigns:
            writer.writerow([
                c.id,
                c.name,
                c.status,
                c.recipients_count,
                c.delivered_count,
                c.read_count,
                c.failed_count,
                f"{c.actual_cost:.2f}",
                f"{c.attributed_revenue:.2f}",
                c.created_at.strftime("%Y-%m-%d %H:%M:%S") if c.created_at else ""
            ])
        return output.getvalue()
